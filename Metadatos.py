# GNU GENERAL PUBLIC LICENSE
# Version 3, 29 June 2007

# Copyright (C) 2024 Moisés Ceñera Fernández

# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.

# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
# GNU General Public License for more details.

# You should have received a copy of the GNU General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.

import wx
import os
from PyPDF2 import PdfReader, PdfWriter
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image
import piexif
import zipfile
import shutil
import xml.etree.ElementTree as ET

class MainApp(wx.App):
    def OnInit(self):
        self.frame = MetadataAnalyzerFrame(None, title="Metadata Analyzer", size=(800, 600))
        self.frame.Show()
        return True

class MetadataAnalyzerFrame(wx.Frame):
    def __init__(self, *args, **kw):
        super(MetadataAnalyzerFrame, self).__init__(*args, **kw)
        
        panel = wx.Panel(self)
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        instructions = wx.StaticText(panel, label="Seleccione una acción para analizar o limpiar los metadatos de archivos.")
        instruction_font = wx.Font(12, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_NORMAL)
        instructions.SetFont(instruction_font)
        main_sizer.Add(instructions, 0, wx.ALL | wx.CENTER, 10)
        
        buttons_sizer = wx.GridBagSizer(10, 10) 
        
        button_font = wx.Font(11, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_BOLD)
        button_size = (250, 50)
        
        select_file_btn = wx.Button(panel, label="📄 Seleccionar Archivo", size=button_size)
        select_file_btn.SetFont(button_font)
        select_file_btn.Bind(wx.EVT_BUTTON, self.on_select_file)
        buttons_sizer.Add(select_file_btn, pos=(0, 0), flag=wx.EXPAND | wx.ALL, border=5)
        
        remove_file_btn = wx.Button(panel, label="🗑️ Eliminar Metadatos (Archivo)", size=button_size)
        remove_file_btn.SetFont(button_font)
        remove_file_btn.Bind(wx.EVT_BUTTON, self.on_remove_metadata_file)
        buttons_sizer.Add(remove_file_btn, pos=(1, 0), flag=wx.EXPAND | wx.ALL, border=5)
        
        select_directory_btn = wx.Button(panel, label="📂 Seleccionar Carpeta", size=button_size)
        select_directory_btn.SetFont(button_font)
        select_directory_btn.Bind(wx.EVT_BUTTON, self.on_select_directory)
        buttons_sizer.Add(select_directory_btn, pos=(0, 1), flag=wx.EXPAND | wx.ALL, border=5)
        
        remove_directory_btn = wx.Button(panel, label="🗑️📂 Eliminar Metadatos (Carpeta)", size=button_size)
        remove_directory_btn.SetFont(button_font)
        remove_directory_btn.Bind(wx.EVT_BUTTON, self.on_remove_metadata_directory)
        buttons_sizer.Add(remove_directory_btn, pos=(1, 1), flag=wx.EXPAND | wx.ALL, border=5)
        
        clear_results_btn = wx.Button(panel, label="🧹 Limpiar Resultados", size=button_size)
        clear_results_btn.SetFont(button_font)
        clear_results_btn.Bind(wx.EVT_BUTTON, self.on_clear_results)
        buttons_sizer.Add(clear_results_btn, pos=(2, 0), span=(1,2), flag=wx.EXPAND | wx.ALL | wx.ALIGN_CENTER_HORIZONTAL, border=5)
        
        buttons_sizer.AddGrowableCol(0, 1)
        buttons_sizer.AddGrowableCol(1, 1)
        
        main_sizer.Add(buttons_sizer, 0, wx.ALL | wx.CENTER, 15)
        
        self.result_text_metadata = wx.TextCtrl(panel, style=wx.TE_MULTILINE | wx.TE_READONLY | wx.HSCROLL, size=(750, 250))
        main_sizer.Add(self.result_text_metadata, 1, wx.ALL | wx.EXPAND, 10)
        
        panel.SetSizer(main_sizer)

    def on_select_file(self, event):
        with wx.FileDialog(self, "Seleccione un archivo", wildcard="Archivos PDF (*.pdf)|*.pdf|Archivos Office (*.docx;*.xlsx;*.pptx)|*.docx;*.xlsx;*.pptx|Archivos imagen (*.jpeg;*.jpg;*.png)|*.jpeg;*.jpg;*.png", style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST) as file_dialog:
            if file_dialog.ShowModal() == wx.ID_CANCEL:
                return
            path = file_dialog.GetPath()
            metadata = analyze_metadata(path)
            self.display_metadata(metadata)

    def on_select_directory(self, event):
        with wx.DirDialog(self, "Seleccione una carpeta", style=wx.DD_DEFAULT_STYLE) as dir_dialog:
            if dir_dialog.ShowModal() == wx.ID_CANCEL:
                return None
            directory_path = dir_dialog.GetPath()
            if not directory_path:
                return None  
            file_list = [os.path.join(directory_path, f) for f in os.listdir(directory_path) if os.path.isfile(os.path.join(directory_path, f))]
            self.display_directory_metadata(analyze_metadata_directory(file_list))
            
    def on_remove_metadata_file(self, event):
        with wx.FileDialog(self, "Seleccione un archivo", wildcard="Archivos PDF (*.pdf)|*.pdf|Archivos Office (*.docx;*.xlsx;*.pptx)|*.docx;*.xlsx;*.pptx|Archivos imagen (*.jpeg;*.jpg;*.png)|*.jpeg;*.jpg;*.png", style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST) as file_dialog:
            if file_dialog.ShowModal() == wx.ID_CANCEL:
                return
            path = file_dialog.GetPath()
            result = remove_metadata_file(path)
            self.display_result(result)

    def on_remove_metadata_directory(self, event):
        with wx.DirDialog(self, "Seleccione una carpeta") as dir_dialog:
            if dir_dialog.ShowModal() == wx.ID_CANCEL:
                return
            directory_path = dir_dialog.GetPath()
            if not directory_path:
                return None  
            file_list = [os.path.join(directory_path, f) for f in os.listdir(directory_path) if os.path.isfile(os.path.join(directory_path, f))]
            results = remove_metadata_directory(file_list)
            self.display_directory_results(results)

    def on_clear_results(self, event):
        self.result_text_metadata.Clear()

    def display_metadata(self, data):
        self.result_text_metadata.Clear()
        if data:
            for key, value in data.items():
                self.result_text_metadata.AppendText(f"{key}: {value}\n")
        else:
            wx.MessageBox("No se encontraron metadatos o el archivo seleccionado es inválido.", "Advertencia", wx.OK | wx.ICON_WARNING)

    def display_result(self, data):
        self.result_text_metadata.Clear()
        if data:
            self.result_text_metadata.AppendText(f"{data}\n")
        else:
            wx.MessageBox("No se pudo eliminar los metadatos o el archivo seleccionado es inválido.", "Advertencia", wx.OK | wx.ICON_WARNING)

    def display_directory_metadata(self, data):
        self.result_text_metadata.Clear()
        if data:
            for file_data in data:
                filename = file_data.get("filename", "Archivo desconocido")
                metadata = file_data.get("metadata", {})
                self.result_text_metadata.AppendText(f"Archivo: {filename}\n")
                if isinstance(metadata, dict):
                    for key, value in metadata.items():
                        self.result_text_metadata.AppendText(f"  {key}: {value}\n")
                else:
                    self.result_text_metadata.AppendText(f"  {metadata}\n")
                self.result_text_metadata.AppendText("\n" + "-" * 40 + "\n\n")
        else:
            wx.MessageBox("No se encontraron metadatos o no se seleccionaron archivos válidos.", "Advertencia", wx.OK | wx.ICON_WARNING)

    def display_directory_results(self, data):
        self.result_text_metadata.Clear()
        if data:
            for message in data:
                self.result_text_metadata.AppendText(f"{message}\n")
        else:
            wx.MessageBox("No se pudo eliminar los metadatos o la selección fue inválida.", "Advertencia", wx.OK | wx.ICON_WARNING)

def analyze_metadata(filepath):
    try:
        if filepath.endswith('.pdf'):
            with open(filepath, 'rb') as file:
                pdf = PdfReader(file)
                info = pdf.metadata
                return info
        elif filepath.endswith('.docx'):
            doc = Document(filepath)
            props = doc.core_properties
            return {
                "Título": props.title or "N/A",
                "Autor": props.author or "N/A",
                "Autor anterior": props.last_modified_by or "N/A",
                "Fecha de creación": props.created or "N/A",
                "Última modificación": props.modified or "N/A",
                "Categoría": props.category or "N/A",
                "Comentarios": props.comments or "N/A"
            }
        elif filepath.endswith('.xlsx'):
            workbook = load_workbook(filepath, read_only=True)
            props = workbook.properties
            if props.creator == "openpyxl":
                props.creator = None
                props.created = None
                props.modified = None
            
            return {
                "Título": props.title or "N/A",
                "Autor": props.creator or "N/A",
                "Autor anterior": props.lastModifiedBy or "N/A",
                "Fecha de creación": props.created or "N/A",
                "Última modificación": props.modified or "N/A",
                "Categoría": props.category or "N/A",
                "Comentarios": props.description or "N/A"
            }
        elif filepath.endswith('.pptx'):
            presentation = Presentation(filepath)
            props = presentation.core_properties
            return {
                "Título": props.title or "N/A",
                "Autor": props.author or "N/A",
                "Autor anterior": props.last_modified_by or "N/A",
                "Fecha de creación": props.created or "N/A",
                "Última modificación": props.modified or "N/A",
                "Categoría": props.category or "N/A",
                "Comentarios": props.comments or "N/A"
            }

        elif filepath.endswith('.jpg') or filepath.endswith('.jpeg') or filepath.endswith('.png'):
            image = Image.open(filepath)
            if "exif" in image.info:
                exif_data = piexif.load(image.info["exif"])
                metadata = {}
                for ifd in exif_data:
                    for tag, value in exif_data[ifd].items():
                        tag_name = piexif.TAGS[ifd].get(tag, tag)
                        metadata[tag_name] = value
                return metadata
    except Exception as e:
        wx.MessageBox(f"Error analizando metadatos: {e}", "Error", wx.OK | wx.ICON_ERROR)

def remove_metadata_pdf(filepath):
    reader = PdfReader(filepath)
    writer = PdfWriter()
    for page in range(len(reader.pages)):
        writer.add_page(reader.pages[page])
    writer.add_metadata({})
    with open(filepath, "wb") as f:
        writer.write(f)

def remove_metadata_office(filepath):
    temp_dir = "temp_file"
    file_extension = filepath.split('.')[-1]
    
    with zipfile.ZipFile(filepath, 'r') as zip_ref:
        zip_ref.extractall(temp_dir)
    
    metadata_files = {
        'docx': ['docProps/core.xml', 'docProps/app.xml'],
        'xlsx': ['docProps/core.xml', 'docProps/app.xml'],
        'pptx': ['docProps/core.xml', 'docProps/app.xml']
    }
    
    if file_extension not in metadata_files:
        raise ValueError("Formato de archivo no soportado para eliminación de metadatos")
    
    for meta_file in metadata_files[file_extension]:
        meta_path = os.path.join(temp_dir, meta_file)
        if os.path.exists(meta_path):
            tree = ET.parse(meta_path)
            root = tree.getroot()
            for elem in root.iter():
                elem.clear()
            tree.write(meta_path)

    with zipfile.ZipFile(filepath, 'w') as zip_ref:
        for folder_name, subfolders, filenames in os.walk(temp_dir):
            for filename in filenames:
                file_path = os.path.join(folder_name, filename)
                zip_ref.write(file_path, os.path.relpath(file_path, temp_dir))

    shutil.rmtree(temp_dir)

def remove_metadata_image(filepath):
    imagen = Image.open(filepath)
    if "exif" in imagen.info:
        imagen.info.pop("exif")
    imagen.save(filepath)

def remove_metadata_file(filepath):
    try:
        extension = os.path.splitext(filepath)[1]
        if extension in ['.jpg', '.jpeg', '.png']:
            remove_metadata_image(filepath)
            return f"Archivo: {os.path.basename(filepath)}  se han borrado los metadatos correctamente."
        elif extension == '.pdf':
            remove_metadata_pdf(filepath)
            return f"Archivo: {os.path.basename(filepath)}  se han borrado los metadatos correctamente."
        elif extension in ['.docx', '.xlsx', '.pptx']:
            remove_metadata_office(filepath)
            return f"Archivo: {os.path.basename(filepath)}  se han borrado los metadatos correctamente."
        else:
            return f"El programa no soporta borrado de metadatos para la extension: {extension} de {os.path.basename(filepath)}"
    except Exception as e:
        wx.MessageBox(f"Error borrando los metadatos: {e}", "Error", wx.OK | wx.ICON_ERROR)

def remove_metadata_directory(file_list):
    try:
        info_list = []
        for file in file_list:
            info_list.append(remove_metadata_file(file))
        return info_list
    except Exception as e:
        wx.MessageBox(f"Error removing metadata from directory: {e}", "Error", wx.OK | wx.ICON_ERROR)

def analyze_metadata_directory(file_list):
    try:
        info_list = []
        for file_read in file_list:
            info = analyze_metadata(file_read)
            info_list.append({
                "filename": file_read,
                "metadata": info
            })
        return info_list
    except Exception as e:
        wx.MessageBox(f"Error analyzing files in directory: {e}", "Error", wx.OK | wx.ICON_ERROR)

def clear_results_metadata(result_text):
    result_text.Clear()
    wx.MessageBox("Text area has been cleared.", "Results cleared", wx.OK | wx.ICON_INFORMATION)


if __name__ == "__main__":
    app = MainApp()
    app.MainLoop()
